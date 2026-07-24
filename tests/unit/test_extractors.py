import datetime as dt
from pathlib import Path

import docx as docx_lib
import openpyxl
import pytest

from sba.modules.indexer.extractors import (
    XLSX_MAX_ROWS_PER_SHEET,
    ExtractError,
    extract,
)


def build_minimal_pdf(text: str) -> bytes:
    """Минимальный однострочный PDF, собранный вручную (без зависимостей)."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"",  # поток содержимого — ниже
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects[3] = (
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_txt_utf8(tmp_path: Path) -> None:
    file = tmp_path / "note.txt"
    file.write_text("Привет, документ", encoding="utf-8")
    blocks = extract(file)
    assert blocks[0].text == "Привет, документ"


def test_txt_cp1251_fallback(tmp_path: Path) -> None:
    file = tmp_path / "old.txt"
    file.write_bytes("Старый виндовый файл".encode("cp1251"))
    assert "Старый виндовый файл" in extract(file)[0].text


def test_markdown_headings_and_frontmatter(tmp_path: Path) -> None:
    file = tmp_path / "note.md"
    file.write_text(
        "---\ntags: [test]\n---\n"
        "Вступление без заголовка.\n"
        "# План экспедиции\nМаршрут по Ладоге.\n"
        "## Бюджет\nСто тысяч рублей.\n",
        encoding="utf-8",
    )
    blocks = extract(file)
    # frontmatter-теги вынесены в отдельный блок вверху, проза — как раньше
    assert [b.locator for b in blocks] == [
        "теги",
        "",
        "раздел «План экспедиции»",
        "раздел «Бюджет»",
    ]
    assert blocks[0].text == "Теги: #test"
    assert "tags" not in blocks[1].text  # тело прозы не содержит frontmatter
    assert blocks[1].text == "Вступление без заголовка."
    assert "Ладоге" in blocks[2].text


def test_markdown_obsidian_tags_and_wikilinks(tmp_path: Path) -> None:
    file = tmp_path / "note.md"
    file.write_text(
        "---\ntags:\n  - экспедиция\n  - ладога\n---\n"
        "# Заметка\n"
        "Идём с #байдарка и #экспедиция к [[Ладожское озеро]].\n"
        "См. также [[Снаряжение|список вещей]] и [[Ладога#Маршрут]].\n",
        encoding="utf-8",
    )
    blocks = extract(file)
    tags_block = next(b for b in blocks if b.locator == "теги")
    # frontmatter + инлайновые, дедуп регистронезависимый, порядок сохранён
    assert tags_block.text == "Теги: #экспедиция #ладога #байдарка"
    links_block = next(b for b in blocks if b.locator == "связи")
    # алиас отброшен (берётся цель), якорь #Раздел отброшен, дедуп по имени
    assert links_block.text == "Связи: Ладожское озеро, Снаряжение, Ладога"


def test_markdown_no_tags_no_meta_blocks(tmp_path: Path) -> None:
    file = tmp_path / "plain.md"
    file.write_text("# Просто заметка\nБез тегов и ссылок.\n", encoding="utf-8")
    blocks = extract(file)
    assert [b.locator for b in blocks] == ["раздел «Просто заметка»"]


def test_xlsx_sheets_and_values(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Бюджет"
    sheet.append(["Статья", "Сумма", "Дата"])
    sheet.append(["Снаряжение", 100000, dt.datetime(2026, 7, 22)])
    sheet.append([None, None, None])  # пустая строка — пропускается
    sheet.append(["Итого", 100000.0, None])
    second = workbook.create_sheet("Пусто")  # лист без данных — без блока
    second["A1"] = None
    file = tmp_path / "plan.xlsx"
    workbook.save(str(file))

    blocks = extract(file)
    assert [b.locator for b in blocks] == ["лист «Бюджет»"]
    text = blocks[0].text
    assert "Статья │ Сумма │ Дата" in text
    assert "Снаряжение │ 100000 │ 2026-07-22" in text  # float→int, дата ISO
    assert "Итого │ 100000" in text
    assert "None" not in text  # пустые ячейки не просачиваются


def test_xlsx_formula_uses_cached_value(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = 2
    sheet["A2"] = 3
    sheet["A3"] = "=A1+A2"
    file = tmp_path / "calc.xlsx"
    workbook.save(str(file))
    # openpyxl без пересчёта: у формулы нет кэша → приходит пусто, а не «=A1+A2»
    text = "\n".join(b.text for b in extract(file))
    assert "=A1+A2" not in text  # текст формулы в поиск не попадает


def test_xlsx_row_cap(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for i in range(XLSX_MAX_ROWS_PER_SHEET + 50):
        sheet.append([f"строка {i}"])
    file = tmp_path / "big.xlsx"
    workbook.save(str(file))
    blocks = extract(file)
    assert "лист обрезан" in blocks[0].text


def test_pdf_pages(tmp_path: Path) -> None:
    file = tmp_path / "doc.pdf"
    file.write_bytes(build_minimal_pdf("Expedition budget 100k"))
    blocks = extract(file)
    assert blocks[0].locator == "стр. 1"
    assert "Expedition budget" in blocks[0].text


def test_docx_headings(tmp_path: Path) -> None:
    document = docx_lib.Document()
    document.add_heading("Договор подряда", level=1)
    document.add_paragraph("Исполнитель обязуется выполнить работы.")
    document.add_paragraph("Срок — до конца года.")
    file = tmp_path / "contract.docx"
    document.save(str(file))
    blocks = extract(file)
    assert blocks[0].locator == "раздел «Договор подряда»"
    assert "Исполнитель" in blocks[0].text


def test_unsupported_extension(tmp_path: Path) -> None:
    file = tmp_path / "img.png"
    file.write_bytes(b"\x89PNG")
    with pytest.raises(ExtractError):
        extract(file)


def test_broken_pdf_raises_extract_error(tmp_path: Path) -> None:
    file = tmp_path / "broken.pdf"
    file.write_bytes("это не PDF вовсе".encode())
    with pytest.raises(ExtractError):
        extract(file)
